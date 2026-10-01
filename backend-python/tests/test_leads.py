"""Phase 26 — Leads CRM: Python twin of tests/leads.test.js.

Public capture forms (contact/corporate/astrology/kundli) gain structured
contact fields; admins get a full pipeline (NEW → CONTACTED → QUALIFIED →
CONVERTED | LOST), filters, assignment, follow-ups, notes, one-click
conversion into a real manual booking, and the leads Excel export.

Capture dedupe (migration 027 twin): repeats merge into the existing lead."""
import json
import time

import pytest

from tests.conftest import admin_login, login

pytestmark = pytest.mark.asyncio

H = lambda tok: {"Authorization": "Bearer " + tok}


async def test_public_capture_and_admin_pipeline(client, db_session):
    """Capture keeps legacy forms working, validates contact when present,
    audits anonymously as 'public'; the admin pipeline covers every move."""
    # legacy kundli-style capture (no contact at all) still works
    r = await client.post("/api/leads", json={"type": "Kundli", "name": "Legacy Enquiry",
                                              "details": "Born 1995-04-11, Jaipur."})
    assert r.status_code == 201, r.text
    assert r.json()["lead"]["status"] == "NEW"

    # structured capture with contact fields
    r = await client.post("/api/leads", json={"type": "Corporate", "name": "Ritu Corporate",
                                              "mobile": "9876500011", "email": "ritu@corp.example",
                                              "service": "Office opening puja", "location": "Gurugram",
                                              "details": "60 attendees, Diwali week."})
    assert r.status_code == 201, r.text
    lead = r.json()["lead"]
    assert lead["mobile"] == "9876500011" and lead["email"] == "ritu@corp.example"
    assert lead["service"] == "Office opening puja"

    # contact channels validated WHEN present; unknown source refused
    assert (await client.post("/api/leads", json={"type": "Contact", "name": "Bad Mobile",
                                                  "mobile": "123", "email": "a@b.example"})).status_code == 400
    assert (await client.post("/api/leads", json={"type": "Spam", "name": "X",
                                                  "mobile": "9876500012"})).status_code == 400

    aa = {"Authorization": "Bearer " + await admin_login(client)}

    # admin manual capture
    r = await client.post("/api/admin/leads", headers=aa,
                          json={"source": "Partner", "name": "Partner Lead",
                                "mobile": "9876500021", "service": "Satyanarayan", "location": "Delhi"})
    assert r.status_code == 201, r.text
    l1 = r.json()["lead"]
    r = await client.post("/api/admin/leads", headers=aa,
                          json={"source": "Other", "name": "To Lose", "mobile": "9876500022"})
    l3 = r.json()["lead"]
    r = await client.post("/api/admin/leads", headers=aa,
                          json={"source": "Other", "name": "Live Lead", "email": "live@example.com"})
    l2 = r.json()["lead"]

    # list + counts + filters
    body = (await client.get("/api/admin/leads", headers=aa)).json()
    assert body["counts"]["NEW"] >= 4
    assert len(body["leads"]) >= 4
    body = (await client.get("/api/admin/leads?status=NEW", headers=aa)).json()
    assert all(x["status"] == "NEW" for x in body["leads"])
    body = (await client.get("/api/admin/leads?q=9876500021", headers=aa)).json()
    assert len(body["leads"]) == 1, "mobile search hits"
    body = (await client.get("/api/admin/leads?source=Partner", headers=aa)).json()
    assert body["leads"] and all(x["type"] == "Partner" for x in body["leads"])

    # lifecycle
    r = await client.post(f"/api/admin/leads/{l1['id']}/status", headers=aa, json={"status": "CONTACTED"})
    assert r.json()["lead"]["status"] == "CONTACTED"
    r = await client.post(f"/api/admin/leads/{l1['id']}/status", headers=aa, json={"status": "QUALIFIED"})
    assert r.json()["lead"]["status"] == "QUALIFIED"
    r = await client.post(f"/api/admin/leads/{l3['id']}/status", headers=aa, json={"status": "LOST"})
    assert r.status_code == 400, "LOST needs a reason"
    r = await client.post(f"/api/admin/leads/{l3['id']}/status", headers=aa,
                          json={"status": "LOST", "reason": "Budget mismatch after three calls."})
    assert r.json()["lead"]["status"] == "LOST"

    # assign + followup + notes
    r = await client.post(f"/api/admin/leads/{l1['id']}/assign", headers=aa, json={"userId": "admin1"})
    assert r.json()["lead"]["assignedTo"] == "admin1"
    r = await client.post(f"/api/admin/leads/{l1['id']}/assign", headers=aa, json={"userId": "u_nobody"})
    assert r.status_code == 400
    r = await client.post(f"/api/admin/leads/{l1['id']}/followup", headers=aa,
                          json={"when": int(time.time() * 1000) + 3600_000})
    assert r.json()["lead"]["followUpAt"] > time.time() * 1000
    r = await client.post(f"/api/admin/leads/{l1['id']}/followup", headers=aa, json={"when": "soon"})
    assert r.status_code == 400
    r = await client.post(f"/api/admin/leads/{l1['id']}/notes", headers=aa, json={"notes": "Wants a morning muhurat."})
    assert "muhurat" in r.json()["lead"]["details"]

    # conversion through the real bookings engine
    r = await client.post(f"/api/admin/leads/{l1['id']}/convert", headers=aa,
                          json={"pujaId": "satyanarayan", "mode": "home", "slot": "10:00 AM"})
    assert r.status_code == 201, r.text
    assert r.json()["bookingId"].startswith("DP"), "a real booking id comes back"
    assert r.json()["lead"]["status"] == "CONVERTED"
    assert r.json()["lead"]["convertedBookingId"] == r.json()["bookingId"]
    assert (await client.post(f"/api/admin/leads/{l1['id']}/convert", headers=aa, json={})).status_code == 409

    # no contact -> convert refuses with guidance
    r = await client.post("/api/admin/leads", headers=aa,
                          json={"source": "Other", "name": "Email Only", "email": "eo@example.com"})
    no_mob = r.json()["lead"]
    assert (await client.post(f"/api/admin/leads/{no_mob['id']}/convert", headers=aa, json={})).status_code == 400

    # delete rules
    assert (await client.delete(f"/api/admin/leads/{l2['id']}", headers=aa)).status_code == 409, "live leads are never deleted"
    assert (await client.delete(f"/api/admin/leads/{l3['id']}", headers=aa)).status_code == 200
    assert (await client.delete(f"/api/admin/leads/{l3['id']}", headers=aa)).status_code == 404

    # audit trail
    from sqlalchemy import select

    from app.models import AuditLog
    rows = (await db_session.execute(
        select(AuditLog).where(AuditLog.action.like("lead.%")).order_by(AuditLog.id.desc()).limit(50))).scalars().all()
    actions = {a.action for a in rows}
    assert {"lead.captured", "lead.status", "lead.assign", "lead.followup",
            "lead.notes", "lead.converted", "lead.deleted"} <= actions
    assert any(a.actor_role == "public" for a in rows if a.action == "lead.captured"), \
        "anonymous captures audit as public"

    # access: admin-only
    ct = await login(client, "customer")
    assert (await client.get("/api/admin/leads", headers={"Authorization": "Bearer " + ct})).status_code == 403
    assert (await client.get("/api/admin/leads")).status_code == 401
    assert (await client.post("/api/admin/leads", headers={"Authorization": "Bearer " + ct},
                              json={"source": "Other", "name": "X"})).status_code == 403


async def test_leads_report_export(client, db_session):
    """The leads Excel export renders with the CRM columns and filters."""
    from openpyxl import load_workbook

    aa = {"Authorization": "Bearer " + await admin_login(client)}
    await client.post("/api/leads", json={"type": "Contact", "name": "Export Probe",
                                          "mobile": "9876500031", "location": "Delhi"})
    r = await client.get("/api/admin/export/leads.xlsx", headers=aa)
    assert r.status_code == 200, r.text
    wb = load_workbook(__import__("io").BytesIO(r.content))
    ws = wb.active
    header = [str(c.value or "") for c in ws[4]]
    assert "Status" in header and "Converted booking" in header, "CRM columns in the export"
    r = await client.get("/api/admin/export/leads.xlsx?status=NEW", headers=aa)
    assert r.status_code == 200


async def test_assignment_ready_conversion(client, db_session):
    """Twin of tests/leads.test.js 'assignment-ready conversion': the picker lists
    availability-filtered pandits, convert with a pandit creates the booking AND
    assigns it, and a second same-slot convert 409s through the engine."""
    import datetime

    from sqlalchemy import select

    from app.models import AuditLog

    aa = {"Authorization": "Bearer " + await admin_login(client)}
    r = await client.post("/api/admin/leads", headers=aa,
                          json={"source": "Walk-in", "name": "Assign Ready",
                                "mobile": "9876500041", "location": "Delhi NCR"})
    lead = r.json()["lead"]

    day = (datetime.date.today() + datetime.timedelta(days=5)).isoformat()
    pick = await client.get("/api/admin/leads/available-pandits",
                            headers=aa,
                            params={"pujaId": "satyanarayan", "mode": "home",
                                    "date": day, "slot": "10:00 AM", "city": "Delhi NCR"})
    assert pick.status_code == 200, pick.text
    assert isinstance(pick.json()["pandits"], list)
    ct = await login(client, "customer")
    assert (await client.get("/api/admin/leads/available-pandits",
                             headers={"Authorization": "Bearer " + ct},
                             params={"date": day, "slot": "10:00 AM"})).status_code == 403
    assert (await client.get("/api/admin/leads/available-pandits",
                             params={"date": day, "slot": "10:00 AM"})).status_code == 401
    if not pick.json()["pandits"]:
        return
    pid = pick.json()["pandits"][0]["id"]

    conv = await client.post(f"/api/admin/leads/{lead['id']}/convert", headers=aa,
                             json={"pujaId": "satyanarayan", "mode": "home",
                                   "slot": "10:00 AM", "date": day, "panditId": pid})
    assert conv.status_code == 201, conv.text
    assert conv.json()["panditId"] == pid
    rows = (await db_session.execute(
        select(AuditLog).where(AuditLog.action == "lead.converted")
        .order_by(AuditLog.id.desc()).limit(5))).scalars().all()
    assert any(pid in (a.detail or "") for a in rows), "conversion audit carries the pandit"

    # same pandit, same slot again -> 409 through the availability engine
    r = await client.post("/api/admin/leads", headers=aa,
                          json={"source": "Other", "name": "Second Slot",
                                "mobile": "9876500042", "location": "Delhi NCR"})
    l2 = r.json()["lead"]
    clash = await client.post(f"/api/admin/leads/{l2['id']}/convert", headers=aa,
                              json={"pujaId": "satyanarayan", "mode": "home",
                                    "slot": "10:00 AM", "date": day, "panditId": pid})
    assert clash.status_code == 409, clash.text
    assert "not available" in clash.json()["detail"]


async def test_capture_dedupe(client, db_session):
    """Twin of tests/leads.test.js 'capture dedupe': a repeat capture whose
    mobile or email matches an existing lead merges into it — no new row,
    fields backfilled, note line kept, dup counter advanced; LOST re-opens,
    CONVERTED is terminal, contact-less enquiries cannot dedupe."""
    from sqlalchemy import func, select

    from app.models import AuditLog, Lead

    # first capture creates; the repeat merges into it
    r = await client.post("/api/leads", json={"type": "Contact", "name": "Dup Probe",
                                              "mobile": "9876510001",
                                              "email": "dupprobe@example.com",
                                              "location": "Varanasi"})
    assert r.status_code == 201, r.text
    first = r.json()["lead"]
    rep = await client.post("/api/leads", json={"type": "Astrology", "name": "Dup Probe Again",
                                                "mobile": "9876510001",
                                                "details": "Called again about navagraha."})
    assert rep.status_code == 201, rep.text
    merged = rep.json()["lead"]
    assert merged["id"] == first["id"], "the repeat returns the EXISTING lead"
    assert merged["merged"] is True, "the response says merged"
    n = (await db_session.execute(select(func.count()).select_from(Lead)
                                  .where(Lead.mobile == "9876510001"))).scalar()
    assert n == 1, "no second row exists"
    assert merged["dupCount"] == 1, "dup counter advanced"
    assert (merged["lastDupAt"] or 0) > 0, "last dup time stamped"
    assert "Called again about navagraha" in merged["details"], "the new message is kept"
    assert "Astrology" in merged["details"], "the note line names the source"

    # email match merges too (case-insensitive); empty fields backfill
    r = await client.post("/api/leads", json={"type": "Corporate", "name": "Dup By Email",
                                              "email": "DUPPROBE@example.com",
                                              "service": "Office puja"})
    rep2 = r.json()["lead"]
    assert rep2["id"] == first["id"], "email match is case-insensitive"
    assert rep2["dupCount"] == 2
    assert rep2["service"] == "Office puja", "empty fields are backfilled from the repeat"
    assert rep2["mobile"] == "9876510001", "existing fields are never overwritten"

    # distinct prospect -> a real new row
    r = await client.post("/api/leads", json={"type": "Kundli", "name": "Fresh Prospect",
                                              "mobile": "9876510002"})
    other = r.json()["lead"]
    assert other["id"] != first["id"]
    assert not other.get("merged")
    assert other["dupCount"] == 0

    # a LOST lead asking again re-opens as NEW (same row, fresh opportunity)
    aa = {"Authorization": "Bearer " + await admin_login(client)}
    r = await client.post(f"/api/admin/leads/{other['id']}/status", headers=aa,
                          json={"status": "LOST", "reason": "Went cold."})
    assert r.json()["lead"]["status"] == "LOST"
    r = await client.post("/api/leads", json={"type": "Contact", "name": "Fresh Prospect",
                                              "mobile": "9876510002",
                                              "details": "Back after the festival."})
    re_ = r.json()["lead"]
    assert re_["id"] == other["id"], "the repeat still merges"
    assert re_["status"] == "NEW", "lost lead re-opened by the new enquiry"
    assert "Back after the festival" in re_["details"]

    # a CONVERTED lead is terminal — a re-enquiry starts a fresh row
    r = await client.post("/api/admin/leads", headers=aa,
                          json={"source": "Walk-in", "name": "Converted Prospect",
                                "mobile": "9876510003"})
    conv_src = r.json()["lead"]
    r = await client.post(f"/api/admin/leads/{conv_src['id']}/convert", headers=aa,
                          json={"pujaId": "satyanarayan", "mode": "home", "slot": "10:00 AM"})
    assert r.status_code == 201, r.text
    r = await client.post("/api/leads", json={"type": "Contact", "name": "Converted Prospect",
                                              "mobile": "9876510003"})
    re_conv = r.json()["lead"]
    assert re_conv["id"] != conv_src["id"], "converted rows never merge"
    assert "merged" not in re_conv
    n = (await db_session.execute(select(func.count()).select_from(Lead)
                                  .where(Lead.mobile == "9876510003"))).scalar()
    assert n == 2, "the new enquiry created its own row"

    # no contact channels -> always a new row (nothing to match on)
    r1 = await client.post("/api/leads", json={"type": "Kundli", "name": "Anon A", "details": "x"})
    r2 = await client.post("/api/leads", json={"type": "Kundli", "name": "Anon A", "details": "y"})
    assert r1.json()["lead"]["id"] != r2.json()["lead"]["id"], "contact-less enquiries cannot dedupe"

    # audits + admin list visibility
    rows = (await db_session.execute(
        select(AuditLog).where(AuditLog.action == "lead.deduped")
        .order_by(AuditLog.id.desc()).limit(20))).scalars().all()
    assert len(rows) >= 3, "lead.deduped audited"
    assert any(a.actor_role == "public" for a in rows), "public merges audit as public"
    assert all(json.loads(a.detail or "{}").get("mode") == "merge" for a in rows)
    body = (await client.get("/api/admin/leads?dupes=1", headers=aa)).json()
    assert body["leads"] and all(x["dupCount"] > 0 for x in body["leads"]), \
        "dupes filter returns only merged rows"
    assert any(x["id"] == first["id"] for x in body["leads"])
